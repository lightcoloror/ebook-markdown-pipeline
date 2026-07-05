from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from ebook_converter_mcp import read_artifact


def load_wrapper_module():
    script = Path(__file__).resolve().parent / "paddleocr_vl_image_to_md.py"
    spec = importlib.util.spec_from_file_location("paddleocr_vl_image_to_md", script)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load wrapper module: {script}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main() -> int:
    wrapper = load_wrapper_module()
    with tempfile.TemporaryDirectory(prefix="paddleocr-vl-wrapper-") as tmp:
        root = Path(tmp)
        (root / "base.md").write_text("正文", encoding="utf-8", newline="\n")
        (root / "result.json").write_text(
            json.dumps(
                {
                    "parsing_res_list": [
                        {"block_label": "text", "block_content": "不是表格"},
                        {"block_label": "table", "block_content": "| A | B |\n| --- | --- |\n| 1 | 2 |"},
                    ]
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
            newline="\n",
        )
        output = root / "out.md"
        wrapper.write_normalized_markdown(root / "base.md", output, root)
        rendered = output.read_text(encoding="utf-8")
        if "真实表格" not in rendered or "| A | B |" not in rendered:
            raise RuntimeError(f"Expected JSON table block to be appended: {rendered}")
        sidecars = wrapper.write_paddleocr_vl_sidecars(root, root / "sample.png", output, rendered, ["paddleocr", "doc_parser"], status="review")
        if {path.name for path in sidecars} != {"document-vlm-result.json", "table-candidates.json"}:
            raise RuntimeError(f"Unexpected PaddleOCR-VL sidecars: {sidecars}")
        document_summary = read_artifact({"path": str(root / "document-vlm-result.json"), "artifact_type": "document_vlm_result_json"}).get("summary") or {}
        if not document_summary.get("schema_valid") or document_summary.get("block_count") != 2 or document_summary.get("table_count") != 1:
            raise RuntimeError(f"Unexpected PaddleOCR-VL document sidecar summary: {document_summary}")
        table_summary = read_artifact({"path": str(root / "table-candidates.json"), "artifact_type": "table_candidates_json"}).get("summary") or {}
        if not table_summary.get("schema_valid") or table_summary.get("table_count") != 1:
            raise RuntimeError(f"Unexpected PaddleOCR-VL table sidecar summary: {table_summary}")

    try:
        from docx import Document
    except Exception:
        print("python-docx unavailable; skipped docx table assertion.")
        return 0

    with tempfile.TemporaryDirectory(prefix="paddleocr-vl-wrapper-docx-") as tmp:
        root = Path(tmp)
        (root / "base.md").write_text("正文", encoding="utf-8", newline="\n")
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "分数"
        table.cell(1, 0).text = "甲"
        table.cell(1, 1).text = "98"
        doc.save(root / "table.docx")
        output = root / "out.md"
        wrapper.write_normalized_markdown(root / "base.md", output, root)
        rendered = output.read_text(encoding="utf-8")
        if "| 姓名 | 分数 |" not in rendered or "| 甲 | 98 |" not in rendered:
            raise RuntimeError(f"Expected DOCX table to be converted: {rendered}")
    print("PaddleOCR-VL wrapper table extraction test passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
